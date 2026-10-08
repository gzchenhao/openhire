"""Write docs/openhire-岗位表.xlsx, the sheet an employer fills in (no IT staff needed)."""
from __future__ import annotations

import pathlib

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = pathlib.Path(__file__).resolve().parent.parent / "docs" / "openhire-岗位表.xlsx"
COLUMNS = ["职位名称", "工作城市", "发布日期", "投递链接或投递邮箱", "薪资范围", "岗位描述"]
WIDTHS = [26, 14, 14, 36, 16, 70]
EXAMPLE = ["感知算法工程师", "上海", "2026-10-08", "hr@example.com",
           "25k-40k/月", "负责多传感器融合与 BEV 感知模型开发。要求：3 年以上感知算法经验，熟悉 PyTorch。"]
NOTES = [
    "怎么填（两分钟）",
    "1. 「岗位」这一页每行一个岗位，六列都用中文填，不用改格式。示例那一行删掉或覆盖都行。",
    "2. 发布日期写岗位真正对外开放的那天，格式 2026-10-08。我们会按它算在架天数，并在数据里标明「雇主自报」。",
    "3. 投递链接要是贵公司自己网站上的页面（https 开头）；没有的话写一个公司邮箱（如 hr@公司域名），不要写个人手机或微信。",
    "4. 薪资范围可以空着。写的话用 25k-40k/月 或 30万-50万/年 这种写法。",
    "5. 岗位描述把职责和要求直接粘进来即可，长短不限。",
    "",
    "怎么发",
    "用贵公司企业邮箱（公司域名的邮箱）把这个文件发到 gdchenhao@qq.com，主题写「岗位表 + 公司名」，正文留一个贵公司官网招聘页的网址。企业邮箱发出来本身就是身份核验，不需要营业执照。",
    "3 个工作日内上线或回信。免费，只读，不改排序，不卖任何东西。",
    "",
    "90 天规则",
    "每条岗位从发布日期起 90 天后自动下架。还在招的，把表再发一次就续上（发布日期不用改，我们按再次收到的日期续）。",
    "不想继续了，发一封邮件说一声就全部下架。",
]


def main() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "岗位"
    ws.append(COLUMNS)
    ws.append(EXAMPLE)
    head_fill = PatternFill("solid", fgColor="E8F0EB")
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True)
        c.fill = head_fill
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    for col in range(1, 7):
        ws.cell(row=2, column=col).font = Font(color="7A7A7A")
    ws.cell(row=2, column=6).alignment = Alignment(wrap_text=True)

    notes = wb.create_sheet("说明")
    notes.column_dimensions["A"].width = 110
    for line in NOTES:
        notes.append([line])
    for row in notes.iter_rows(min_row=1, max_row=len(NOTES)):
        cell = row[0]
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if cell.value in ("怎么填（两分钟）", "怎么发", "90 天规则"):
            cell.font = Font(bold=True)
    wb.save(OUT)
    print("wrote", OUT, OUT.stat().st_size, "bytes")


if __name__ == "__main__":
    main()
